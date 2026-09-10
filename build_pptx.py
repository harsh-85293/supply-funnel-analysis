"""
Turn deliverables/DECK.md into a native PowerPoint: deliverables/DECK.pptx.

Same source of truth as DECK.pdf, so the figures stay exactly the ones
verify_claims.py checks. Six 16:9 slides, charts embedded, count asserted.

    python build_pptx.py
"""
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent
DELIV = ROOT / "deliverables"

# 16:9 canvas
EMU = 914400
SLIDE_W = int(13.333 * EMU)
SLIDE_H = int(7.5 * EMU)
MARGIN = int(0.55 * EMU)
CONTENT_W = SLIDE_W - 2 * MARGIN

INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
BAND = RGBColor(0xF2, 0xF2, 0xF2)
RULE = RGBColor(0xD5, 0xD5, 0xD5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ---------------------------------------------------------------- markdown parse
def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def parse_slides(md):
    """Return a list of slides; each slide is a list of (kind, payload) blocks."""
    lines = md.split("\n")
    slides, cur, i = [], None, 0
    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("## Slide"):
            if cur is not None:
                slides.append(cur)
            cur = [("title", s[3:])]
            i += 1
            continue
        if cur is None or not s or re.fullmatch(r"-{3,}", s):
            i += 1
            continue
        if s.startswith("### "):
            cur.append(("subtitle", s[4:]))
            i += 1
        elif s.startswith("!["):
            m = re.search(r"\((.+?)\)", s)
            cur.append(("image", m.group(1) if m else ""))
            i += 1
        elif s.startswith("|"):
            rows, header = [], None
            while i < len(lines) and lines[i].strip().startswith("|"):
                if is_sep(lines[i]):
                    header = rows.pop() if rows else None
                else:
                    rows.append(split_row(lines[i]))
                i += 1
            cur.append(("table", (header, rows)))
        elif re.match(r"^\d+\.\s+", s) or re.match(r"^[-*]\s+", s):
            items = []
            while i < len(lines):
                c = lines[i].strip()
                if re.match(r"^\d+\.\s+", c):
                    items.append(re.sub(r"^\d+\.\s+", "", c))
                elif re.match(r"^[-*]\s+", c):
                    items.append(re.sub(r"^[-*]\s+", "", c))
                elif c and not c.startswith(("|", "#", "!")) and items:
                    items[-1] += " " + c
                else:
                    break
                i += 1
            cur.append(("bullet", items))
        else:
            para = [s]
            i += 1
            while i < len(lines):
                c = lines[i].strip()
                if not c or c.startswith(("|", "#", "!", "- ", "* ")) or re.fullmatch(r"-{3,}", c):
                    break
                para.append(c)
                i += 1
            cur.append(("para", " ".join(para)))
    if cur is not None:
        slides.append(cur)
    return slides


# ---------------------------------------------------------------- rich runs (**bold**, *italic*)
def add_runs(paragraph, text, size, color=INK, base_bold=False):
    text = (text.replace("—", "\u2014").replace("→", "\u2192")
                .replace("←", "\u2190").replace("·", "\u2022"))
    # split on ** and * while keeping delimiters
    tokens = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text)
    for tok in tokens:
        if not tok:
            continue
        bold, italic, body = base_bold, False, tok
        if tok.startswith("**") and tok.endswith("**"):
            bold, body = True, tok[2:-2]
        elif tok.startswith("*") and tok.endswith("*"):
            italic, body = True, tok[1:-1]
        run = paragraph.add_run()
        run.text = body
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = "Calibri"


# ---------------------------------------------------------------- slide builders
def add_textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(int(0.05 * EMU))
    tf.margin_top = tf.margin_bottom = Emu(int(0.03 * EMU))
    return tb, tf


def add_title(slide, kicker, subtitle, top):
    tb, tf = add_textbox(slide, MARGIN, top, CONTENT_W, int(0.4 * EMU))
    p = tf.paragraphs[0]
    add_runs(p, kicker.upper(), 11, ACCENT, base_bold=True)
    cursor = top + int(0.42 * EMU)
    if subtitle:
        tb2, tf2 = add_textbox(slide, MARGIN, cursor, CONTENT_W, int(1.1 * EMU))
        add_runs(tf2.paragraphs[0], subtitle, 19, INK, base_bold=True)
        # crude line-count estimate for spacing
        lines = max(1, int(len(subtitle) / 68) + 1)
        cursor += int((0.36 * lines + 0.12) * EMU)
    return cursor


def col_widths(raw_rows, ncol, total):
    lengths = []
    for c in range(ncol):
        cells = [len(re.sub(r"[*`]", "", r[c])) for r in raw_rows if c < len(r)] or [1]
        lengths.append(max(1.0, (max(cells) + sum(cells) / len(cells)) / 2))
    lo = total * (0.05 if ncol > 3 else 0.12)
    w = [total * L / sum(lengths) for L in lengths]
    w = [max(x, lo) for x in w]
    k = total / sum(w)
    return [int(x * k) for x in w]


def add_table(slide, header, rows, top, height, font=11):
    data = ([header] if header else []) + rows
    ncol = max(len(r) for r in data)
    data = [r + [""] * (ncol - len(r)) for r in data]
    widths = col_widths(data, ncol, CONTENT_W)

    gtbl = slide.shapes.add_table(len(data), ncol, MARGIN, top, CONTENT_W, height)
    tbl = gtbl.table
    tbl.first_row = bool(header)
    tbl.horz_banding = False
    for c in range(ncol):
        tbl.columns[c].width = widths[c]

    for r in range(len(data)):
        for c in range(ncol):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Emu(int(0.06 * EMU))
            cell.margin_top = cell.margin_bottom = Emu(int(0.02 * EMU))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            is_head = header and r == 0
            cell.fill.fore_color.rgb = BAND if is_head else WHITE
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT
            add_runs(para, data[r][c], font, INK, base_bold=is_head)
    return gtbl


def add_bullets(slide, items, top, height, font=13):
    tb, tf = add_textbox(slide, MARGIN, top, CONTENT_W, height)
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        add_runs(p, "\u2022  " + it, font, INK)


def add_para(slide, text, top, height, font=12, color=INK):
    tb, tf = add_textbox(slide, MARGIN, top, CONTENT_W, height)
    add_runs(tf.paragraphs[0], text, font, color)


def add_image(slide, path, top, max_h, max_w_frac=1.0):
    p = ROOT / path
    if not p.exists():
        raise FileNotFoundError(f"chart missing: {p} (run run_all.py first)")
    from PIL import Image
    iw, ih = Image.open(p).size
    max_w = int(CONTENT_W * max_w_frac)
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    left = MARGIN + (CONTENT_W - w) // 2
    slide.shapes.add_picture(str(p), left, top, w, h)
    return h


def footer(slide, n):
    tb, tf = add_textbox(slide, MARGIN, SLIDE_H - int(0.42 * EMU),
                         CONTENT_W, int(0.3 * EMU))
    p = tf.paragraphs[0]
    add_runs(p, "Captain Acquisition, Supply  \u2014  30 June 2026", 9, GREY)
    p.alignment = PP_ALIGN.LEFT
    tb2, tf2 = add_textbox(slide, SLIDE_W - MARGIN - int(1.2 * EMU),
                           SLIDE_H - int(0.42 * EMU), int(1.1 * EMU), int(0.3 * EMU))
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    add_runs(p2, f"{n} / 6", 9, GREY)


# ---------------------------------------------------------------- build
def _estimate_heights(blocks, top, sc):
    """Estimate the height each body block will occupy at scale `sc` (in EMU)."""
    body = [(k, p) for k, p in blocks if k in ("table", "bullet", "para", "image")]
    from PIL import Image
    # if the slide also carries a table, the chart must yield room for it
    has_table = any(k == "table" for k, _ in body)
    heights = []
    for k, payload in body:
        if k == "image":
            p = ROOT / payload
            iw, ih = Image.open(p).size
            frac = 0.38 if has_table else 0.60
            budget_h = (SLIDE_H - top - int(0.55 * EMU)) * frac
            scale = min(CONTENT_W / iw, budget_h / ih)
            heights.append((k, payload, int(ih * scale) + int(0.14 * EMU)))
        elif k == "table":
            header, rows = payload
            nrows = len(rows) + (1 if header else 0)
            # PowerPoint enforces a minimum row height (~0.4"); estimate a touch
            # generously so a following paragraph never lands on the last row
            row_h = max(0.40, 0.40 * sc) * EMU
            heights.append((k, payload, int(nrows * row_h + 0.28 * EMU)))
        elif k == "bullet":
            per = 0.40 * sc
            hh = sum(max(1, int(len(it) / (95 / max(sc, .6))) + 1) for it in payload) * per
            heights.append((k, payload, int((hh + 0.12) * EMU)))
        else:  # para
            lines = max(1, int(len(payload) / (95 / max(sc, .6))) + 1)
            heights.append((k, payload, int((lines * 0.26 * sc + 0.16) * EMU)))
    return heights


def build(md_path, out):
    slides = parse_slides(md_path.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    bottom_limit = SLIDE_H - int(0.55 * EMU)

    for n, blocks in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        kicker = next((p for k, p in blocks if k == "title"), "")
        num, _, rest = kicker.partition("\u2014")
        if "\u2014" not in kicker:
            num, rest = kicker, ""
        subtitle = next((p for k, p in blocks if k == "subtitle"), "")
        top0 = add_title(slide, num.strip(), (rest.strip() or subtitle), int(0.45 * EMU))
        has_image = any(k == "image" for k, _ in blocks)

        # choose the largest content scale that fits vertically, reserving the
        # footer band plus a safety margin so no text box can land on the footer
        sc = 1.0
        safe_limit = bottom_limit - int(0.22 * EMU)
        for _ in range(12):
            est = _estimate_heights(blocks, top0, sc)
            gaps = int(0.10 * EMU) * len(est)
            if top0 + sum(h for _, _, h in est) + gaps <= safe_limit or sc <= 0.58:
                break
            sc -= 0.05

        base_table = 11 if sc > 0.9 else (10 if sc > 0.78 else 9)
        base_bullet = (13 if not has_image else 12)
        base_bullet = base_bullet if sc > 0.9 else int(base_bullet * 0.9)
        base_para = (12 if not has_image else 11)
        base_para = base_para if sc > 0.9 else int(base_para * 0.9)

        # hard ceiling for any body content: the top of the footer band
        content_ceiling = SLIDE_H - int(0.5 * EMU)
        top = top0
        for k, payload, h in _estimate_heights(blocks, top0, sc):
            # never let a block start inside the footer
            if top >= content_ceiling:
                break
            if k == "image":
                got = add_image(slide, payload, top, h - int(0.14 * EMU),
                                max_w_frac=0.82 if any(kk == "table" for kk, _, _ in
                                                       _estimate_heights(blocks, top0, sc))
                                else 1.0)
                top += got + int(0.16 * EMU)
            elif k == "table":
                header, rows = payload
                nrows = len(rows) + (1 if header else 0)
                font = base_table if nrows <= 6 else max(9, base_table - 1)
                add_table(slide, header, rows, top, h - int(0.1 * EMU), font=font)
                top += h + int(0.06 * EMU)
            elif k == "bullet":
                add_bullets(slide, payload, top, h, font=base_bullet)
                top += h + int(0.06 * EMU)
            else:
                # clamp a closing paragraph so it never rides onto the footer
                avail = content_ceiling - top
                add_para(slide, payload, top, min(h, max(avail, int(0.2 * EMU))),
                         font=base_para)
                top += h + int(0.04 * EMU)

        footer(slide, n)

    prs.save(str(out))
    return out, len(slides)


def main():
    if not (ROOT / "outputs" / "charts").exists():
        print("outputs/charts missing - run `python run_all.py` first.")
        return 1
    out, n = build(DELIV / "DECK.md", DELIV / "DECK.pptx")
    md_slides = len(re.findall(r"^## Slide", (DELIV / "DECK.md").read_text(encoding="utf-8"), re.M))
    print(f"  DECK.pptx   {n} slide(s)   limit 6   {'OK' if n <= 6 else 'OVER LIMIT'}")
    print(f"  slides parsed from DECK.md: {md_slides}")
    if n > 6 or n != md_slides:
        print("\nFAIL: slide count wrong.")
        return 1
    print(f"\nWrote {out.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
